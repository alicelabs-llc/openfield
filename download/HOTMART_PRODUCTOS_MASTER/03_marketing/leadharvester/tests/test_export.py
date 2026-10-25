"""Tests for export functionality."""
import pytest
import json
import csv
from datetime import datetime
from pathlib import Path

from leadharvester.core.models import Lead, LeadList, LeadSource
from leadharvester.export import to_json, to_csv, to_excel


@pytest.fixture
def sample_leads():
    return [
        Lead(
            username="tienda_ejemplo",
            full_name="Tienda Ejemplo SAC",
            bio="Ventas de ropa y accesorios. DM para pedidos. hola@tiendaejemplo.com",
            followers=12500,
            following=240,
            posts_count=480,
            is_business=True,
            category="moda",
            email="hola@tiendaejemplo.com",
            phone="+593 99 123 4567",
            website="https://tiendaejemplo.com",
            source=LeadSource.INSTAGRAM,
        ),
        Lead(
            username="user_normal",
            full_name="User Normal",
            bio="Living life 🌴",
            followers=350,
            following=500,
            posts_count=20,
            is_business=False,
            source=LeadSource.INSTAGRAM,
        ),
    ]


@pytest.fixture
def lead_list(sample_leads):
    return LeadList(name="test", query="test_query", leads=sample_leads)


def test_to_json(lead_list, tmp_path):
    path = to_json(lead_list, tmp_path / "leads.json")
    assert path.exists()
    data = json.loads(path.read_text(encoding="utf-8"))
    assert data["name"] == "test"
    assert len(data["leads"]) == 2
    assert data["leads"][0]["username"] == "tienda_ejemplo"


def test_to_csv(lead_list, tmp_path):
    path = to_csv(lead_list, tmp_path / "leads.csv")
    assert path.exists()
    with path.open() as f:
        reader = csv.DictReader(f)
        rows = list(reader)
    assert len(rows) == 2
    assert rows[0]["username"] == "tienda_ejemplo"
    assert rows[0]["email"] == "hola@tiendaejemplo.com"
    assert rows[1]["username"] == "user_normal"


def test_to_excel(lead_list, tmp_path):
    from openpyxl import load_workbook
    path = to_excel(lead_list, tmp_path / "leads.xlsx")
    assert path.exists()
    wb = load_workbook(str(path))
    assert "Leads" in wb.sheetnames
    assert "Summary" in wb.sheetnames
    ws = wb["Leads"]
    assert ws.cell(row=1, column=1).value == "Username"
    assert ws.cell(row=2, column=1).value == "tienda_ejemplo"
    ws2 = wb["Summary"]
    assert ws2.cell(row=2, column=2).value == 2  # Total leads


def test_leadlist_aggregates(lead_list):
    assert lead_list.count == 2
    assert lead_list.with_email == 1
    assert lead_list.business_count == 1
