"""LeadHarvester — Export to CSV/Excel/JSON."""
from __future__ import annotations
import csv
import json
from pathlib import Path
from leadharvester.core.models import LeadList


def to_json(lead_list: LeadList, path: str | Path) -> Path:
    """Export leads to JSON."""
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(lead_list.model_dump_json(indent=2), encoding="utf-8")
    return p


def to_csv(lead_list: LeadList, path: str | Path) -> Path:
    """Export leads to CSV."""
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    fields = [
        "username", "full_name", "bio", "followers", "following", "posts_count",
        "verified", "is_business", "category", "email", "phone", "website",
        "location", "profile_url", "source", "scraped_at",
    ]
    with p.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for lead in lead_list.leads:
            row = lead.model_dump()
            row["source"] = row["source"].value if hasattr(row["source"], "value") else str(row["source"])
            row["scraped_at"] = row["scraped_at"].isoformat() if hasattr(row["scraped_at"], "isoformat") else str(row["scraped_at"])
            writer.writerow({k: row.get(k, "") for k in fields})
    return p


def to_excel(lead_list: LeadList, path: str | Path) -> Path:
    """Export leads to Excel (.xlsx)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    # Sheet 1: Leads
    ws = wb.active
    ws.title = "Leads"
    headers = [
        "Username", "Full Name", "Bio", "Followers", "Following", "Posts",
        "Verified", "Business", "Category", "Email", "Phone", "Website",
        "Location", "Profile URL", "Source", "Scraped At",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="26314F", end_color="26314F", fill_type="solid")

    for lead in lead_list.leads:
        ws.append([
            lead.username, lead.full_name, lead.bio[:200],
            lead.followers, lead.following, lead.posts_count,
            "✓" if lead.verified else "",
            "✓" if lead.is_business else "",
            lead.category, lead.email, lead.phone, lead.website,
            lead.location, lead.profile_url,
            lead.source.value if hasattr(lead.source, "value") else str(lead.source),
            lead.scraped_at.isoformat(),
        ])

    # Auto-width
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

    # Sheet 2: Summary
    ws2 = wb.create_sheet("Summary")
    ws2.append(["Metric", "Value"])
    ws2.append(["Total leads", lead_list.count])
    ws2.append(["With email", lead_list.with_email])
    ws2.append(["With phone", lead_list.with_phone])
    ws2.append(["Business accounts", lead_list.business_count])
    ws2.append(["Verified accounts", sum(1 for l in lead_list.leads if l.verified)])
    ws2.append(["Total followers (sum)", sum(l.followers for l in lead_list.leads)])
    ws2.append(["Avg followers", sum(l.followers for l in lead_list.leads) // max(len(lead_list.leads), 1)])

    wb.save(str(p))
    return p
